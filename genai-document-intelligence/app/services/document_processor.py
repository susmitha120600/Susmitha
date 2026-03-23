import os
import uuid
from typing import List, Dict, Any, Optional
from pathlib import Path
import pypdf
from docx import Document as DocxDocument
from langchain.text_splitter import RecursiveCharacterTextSplitter
from app.core.config import settings
from app.core.logging import get_logger
from app.core.exceptions import DocumentProcessingError
from app.models.schemas import DocumentChunk, ChunkMetadata, DocumentType
from app.utils.text_utils import clean_text

logger = get_logger(__name__)


class DocumentProcessor:
    
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
        logger.info("DocumentProcessor initialized")
    
    async def process_document(
        self,
        file_path: str,
        document_type: DocumentType,
        metadata: Optional[Dict[str, Any]] = None
    ) -> List[DocumentChunk]:
        try:
            logger.info(f"Processing document: {file_path}, type: {document_type}")
            
            text = await self._extract_text(file_path, document_type)
            
            if not text or len(text.strip()) < 10:
                raise DocumentProcessingError("Extracted text is empty or too short")
            
            chunks = await self._chunk_text(text, file_path, metadata or {})
            
            logger.info(f"Document processed successfully: {len(chunks)} chunks created")
            return chunks
            
        except Exception as e:
            logger.error(f"Error processing document {file_path}: {str(e)}")
            raise DocumentProcessingError(f"Failed to process document: {str(e)}")
    
    async def _extract_text(self, file_path: str, document_type: DocumentType) -> str:
        try:
            if document_type == DocumentType.PDF:
                return await self._extract_from_pdf(file_path)
            elif document_type == DocumentType.DOCX:
                return await self._extract_from_docx(file_path)
            elif document_type == DocumentType.TXT:
                return await self._extract_from_txt(file_path)
            elif document_type == DocumentType.PPTX:
                return await self._extract_from_pptx(file_path)
            else:
                raise DocumentProcessingError(f"Unsupported document type: {document_type}")
        except Exception as e:
            logger.error(f"Text extraction error: {str(e)}")
            raise DocumentProcessingError(f"Failed to extract text: {str(e)}")
    
    async def _extract_from_pdf(self, file_path: str) -> str:
        try:
            text_parts = []
            with open(file_path, 'rb') as file:
                pdf_reader = pypdf.PdfReader(file)
                for page_num, page in enumerate(pdf_reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"[Page {page_num + 1}]\n{page_text}")
            
            full_text = "\n\n".join(text_parts)
            return clean_text(full_text)
        except Exception as e:
            raise DocumentProcessingError(f"PDF extraction failed: {str(e)}")
    
    async def _extract_from_docx(self, file_path: str) -> str:
        try:
            doc = DocxDocument(file_path)
            text_parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            full_text = "\n\n".join(text_parts)
            return clean_text(full_text)
        except Exception as e:
            raise DocumentProcessingError(f"DOCX extraction failed: {str(e)}")
    
    async def _extract_from_txt(self, file_path: str) -> str:
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            return clean_text(text)
        except Exception as e:
            raise DocumentProcessingError(f"TXT extraction failed: {str(e)}")
    
    async def _extract_from_pptx(self, file_path: str) -> str:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            text_parts = []
            
            for slide_num, slide in enumerate(prs.slides):
                slide_text = f"[Slide {slide_num + 1}]\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text += shape.text + "\n"
                text_parts.append(slide_text)
            
            full_text = "\n\n".join(text_parts)
            return clean_text(full_text)
        except Exception as e:
            raise DocumentProcessingError(f"PPTX extraction failed: {str(e)}")
    
    async def _chunk_text(
        self,
        text: str,
        source: str,
        metadata: Dict[str, Any]
    ) -> List[DocumentChunk]:
        try:
            document_id = str(uuid.uuid4())
            
            text_chunks = self.text_splitter.split_text(text)
            
            chunks = []
            for idx, chunk_text in enumerate(text_chunks):
                if len(chunk_text.strip()) < settings.MIN_CHUNK_SIZE:
                    continue
                
                chunk_id = f"{document_id}_chunk_{idx}"
                
                chunk_metadata = ChunkMetadata(
                    document_id=document_id,
                    chunk_id=chunk_id,
                    source=source,
                    chunk_index=idx,
                    total_chunks=len(text_chunks),
                    **metadata
                )
                
                chunk = DocumentChunk(
                    chunk_id=chunk_id,
                    text=chunk_text,
                    metadata=chunk_metadata
                )
                
                chunks.append(chunk)
            
            logger.info(f"Created {len(chunks)} chunks from document {document_id}")
            return chunks
            
        except Exception as e:
            logger.error(f"Chunking error: {str(e)}")
            raise DocumentProcessingError(f"Failed to chunk text: {str(e)}")
    
    def validate_file(self, file_path: str, max_size_mb: int = 50) -> bool:
        if not os.path.exists(file_path):
            raise DocumentProcessingError(f"File not found: {file_path}")
        
        file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
        if file_size_mb > max_size_mb:
            raise DocumentProcessingError(
                f"File size ({file_size_mb:.2f}MB) exceeds maximum allowed size ({max_size_mb}MB)"
            )
        
        return True
